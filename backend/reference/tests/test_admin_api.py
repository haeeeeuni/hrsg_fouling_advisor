"""참조 데이터 관리자 API (specs/06 AC-06-1~6, specs/08 AC-08-3·5)."""

import io
from datetime import date, timedelta

import pytest
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from common.models import AuditLog
from reference.models import CalcParameterSet, CleaningMethod, GtModel, SmpPrice
from reference.param_defs import PARAM_DEFS, default_params
from reference.services import GT_COLUMNS

pytestmark = pytest.mark.django_db

GT = "/api/admin/gt-models/"
METHODS = "/api/admin/cleaning-methods/"
SMP = "/api/admin/smp-prices/"
PARAMS = "/api/admin/calc-parameter-sets/"


@pytest.fixture
def as_admin(api, admin_user, seeded):
    api.force_authenticate(admin_user)
    return api


def gt_payload(**overrides):
    data = {
        "name": "실제 모델 X",
        "manufacturer": "제조사",
        "rated_gt_mw": 300,
        "rated_st_mw": 150,
        "design_backpressure_kpa": 3.0,
        "backpressure_alarm_kpa": 4.5,
        "backpressure_trip_kpa": 5.5,
        "design_exhaust_temp_c": 600,
        "exhaust_temp_alarm_c": 630,
        "exhaust_temp_trip_c": 650,
        "design_stack_temp_c": 95,
        "is_placeholder": False,
    }
    data.update(overrides)
    return data


# --- 시드 (AC-06-1) ---


def test_seed_creates_placeholder_reference_data_once(seeded):
    call_command("seed_defaults", verbosity=0)

    assert GtModel.objects.count() == 3
    assert CleaningMethod.objects.count() == 3
    assert CalcParameterSet.objects.count() == 1
    assert SmpPrice.objects.count() == 1
    assert all(m.is_placeholder for m in GtModel.objects.all())
    active = CalcParameterSet.objects.get(is_active=True)
    assert active.version_label == "p1"
    assert active.is_seed
    assert active.params == default_params()
    assert SmpPrice.objects.get().is_estimate


def test_seed_does_not_overwrite_admin_edits(seeded):
    GtModel.objects.filter(name="가상 모델 A (F급)").update(rated_gt_mw=999)

    call_command("seed_defaults", verbosity=0)

    assert GtModel.objects.get(name="가상 모델 A (F급)").rated_gt_mw == 999


# --- GT 한계표 ---


def test_create_gt_model(as_admin):
    res = as_admin.post(GT, gt_payload(), format="json")

    assert res.status_code == 201
    assert res.data["version"] == 1
    assert AuditLog.objects.filter(target_type="GtModel", action="CREATE").exists()


@pytest.mark.parametrize(
    "overrides",
    [
        {"backpressure_alarm_kpa": 5.5, "backpressure_trip_kpa": 5.5},
        {"exhaust_temp_alarm_c": 700, "exhaust_temp_trip_c": 650},
        {"rated_gt_mw": 0},
    ],
)
def test_alarm_must_be_below_trip(as_admin, overrides):
    """AC-06-2"""
    res = as_admin.post(GT, gt_payload(**overrides), format="json")

    assert res.status_code == 400


def test_update_bumps_version_and_detects_concurrent_edit(as_admin):
    """AC-08-5 — 먼저 저장한 쪽이 이기고, 옛 버전으로 저장하면 409."""
    model = GtModel.objects.first()

    first = as_admin.patch(f"{GT}{model.id}/", {"rated_gt_mw": 280, "version": 1}, format="json")
    assert first.status_code == 200
    assert first.data["version"] == 2

    stale = as_admin.patch(f"{GT}{model.id}/", {"rated_gt_mw": 290, "version": 1}, format="json")
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "STALE_VERSION"
    model.refresh_from_db()
    assert model.rated_gt_mw == 280


def test_patch_with_partial_limits_still_validates_against_saved_values(as_admin):
    model = GtModel.objects.get(name="가상 모델 A (F급)")  # 경보 4.5, 트립 5.5

    res = as_admin.patch(f"{GT}{model.id}/", {"backpressure_alarm_kpa": 6.0}, format="json")

    assert res.status_code == 400


def test_delete_deactivates_instead(as_admin):
    model = GtModel.objects.first()

    assert as_admin.delete(f"{GT}{model.id}/").status_code == 204

    model.refresh_from_db()
    assert model.is_active is False
    assert AuditLog.objects.filter(
        target_type="GtModel", action="UPDATE", after={"is_active": False}
    ).exists()
    assert not AuditLog.objects.filter(target_type="GtModel", action="DELETE").exists()


# --- xlsx 가져오기 (AC-06-6) ---


def xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(GT_COLUMNS))
    for row in rows:
        sheet.append([row.get(c) for c in GT_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "gt.xlsx"
    return buffer


def test_import_creates_and_updates(as_admin):
    existing = GtModel.objects.get(name="가상 모델 A (F급)")
    rows = [gt_payload(name="신규 모델"), gt_payload(name=existing.name, rated_gt_mw=275)]

    res = as_admin.post(f"{GT}import/", {"file": xlsx(rows)}, format="multipart")

    assert res.status_code == 200, res.content
    assert res.data == {"created": 1, "updated": 1}
    existing.refresh_from_db()
    assert existing.rated_gt_mw == 275
    assert existing.version == 2


def test_import_is_all_or_nothing(as_admin):
    rows = [gt_payload(name="정상 행"), gt_payload(name="오류 행", backpressure_trip_kpa=1)]

    res = as_admin.post(f"{GT}import/", {"file": xlsx(rows)}, format="multipart")

    assert res.status_code == 400
    assert "3행" in res.json()["error"]["details"]
    assert not GtModel.objects.filter(name="정상 행").exists()


def test_import_rejects_missing_columns_and_non_xlsx(as_admin):
    workbook = Workbook()
    workbook.active.append(["name", "rated_gt_mw"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "partial.xlsx"
    assert (
        as_admin.post(f"{GT}import/", {"file": buffer}, format="multipart").json()["error"]["code"]
        == "MISSING_COLUMNS"
    )

    text = io.BytesIO(b"a,b")
    text.name = "data.csv"
    assert as_admin.post(f"{GT}import/", {"file": text}, format="multipart").status_code == 400


def test_template_round_trips(as_admin):
    res = as_admin.get(f"{GT}template/")

    sheet = load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert list(rows[0]) == list(GT_COLUMNS)
    assert len(rows) == 1 + GtModel.objects.count()


# --- 세정 공법 ---


@pytest.mark.parametrize(
    "overrides", [{"recovery_ratio": 0}, {"recovery_ratio": 1.2}, {"cleaning_cost_won": -1}]
)
def test_method_validation(as_admin, overrides):
    payload = {
        "name": "새 공법",
        "cleaning_cost_won": 1e7,
        "outage_days": 1,
        "recovery_ratio": 0.8,
        **overrides,
    }

    assert as_admin.post(METHODS, payload, format="json").status_code == 400


# --- SMP ---


def test_smp_history_latest_first(as_admin):
    # 시드 SMP 의 기준일이 오늘이라, 그보다 뒤 날짜로 등록한다.
    later, latest = (date.today() + timedelta(days=d) for d in (1, 2))
    as_admin.post(
        SMP, {"value_won_per_kwh": 140, "as_of_date": str(later), "source": "KPX"}, format="json"
    )
    as_admin.post(
        SMP, {"value_won_per_kwh": 160, "as_of_date": str(latest), "source": "KPX"}, format="json"
    )

    rows = as_admin.get(SMP).data["results"]

    assert rows[0]["value_won_per_kwh"] == 160
    assert rows[0]["created_by_name"] == "김관리"


# --- 파라미터 세트 (AC-06-3) ---


def test_parameter_list_includes_definitions(as_admin):
    data = as_admin.get(PARAMS).data

    assert [d["key"] for d in data["definitions"]] == [d.key for d in PARAM_DEFS]
    assert data["results"][0]["version_label"] == "p1"


def test_new_version_activates_and_old_one_can_be_restored(as_admin):
    params = {**default_params(), "capacity_factor": 0.9}

    created = as_admin.post(PARAMS, {"params": params, "note": "실적 반영"}, format="json")
    assert created.status_code == 201
    assert created.data["version_label"] == "p2"
    assert created.data["is_seed"] is False
    assert CalcParameterSet.objects.get(is_active=True).version_label == "p2"

    p1 = CalcParameterSet.objects.get(version_label="p1")
    assert as_admin.post(f"{PARAMS}{p1.id}/activate/").status_code == 200
    assert CalcParameterSet.objects.get(is_active=True).version_label == "p1"
    assert CalcParameterSet.objects.filter(is_active=True).count() == 1
    assert AuditLog.objects.filter(target_type="CalcParameterSet", action="ACTIVATE").exists()


@pytest.mark.parametrize(
    "mutate",
    [
        lambda p: p.pop("capacity_factor"),
        lambda p: p.update(capacity_factor=1.5),
        lambda p: p.update(calc_pinch_range_c=[15, 5]),
        lambda p: p.update(unknown_key=1),
        lambda p: p.update(fuel_cost_ratio="많이"),
    ],
)
def test_invalid_parameters_are_rejected(as_admin, mutate):
    params = default_params()
    mutate(params)

    res = as_admin.post(PARAMS, {"params": params}, format="json")

    assert res.status_code == 400
    assert CalcParameterSet.objects.count() == 1


def test_preview_compares_without_saving(as_admin):
    params = {**default_params(), "dp_power_loss_coeff": 0.7}

    data = as_admin.post(f"{PARAMS}preview/", {"params": params}, format="json").data

    assert data["current"]["param_version"] == "p1"
    assert data["candidate"]["power_loss_gt_mw"] == pytest.approx(
        2 * data["current"]["power_loss_gt_mw"]
    )
    assert CalcParameterSet.objects.count() == 1
