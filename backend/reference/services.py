"""참조 데이터 도메인 로직 — 파라미터 버전, 시드, GT 표 가져오기 (specs/06)."""

from __future__ import annotations

import io
import re
from datetime import date
from typing import Any

from django.db import IntegrityError, transaction
from openpyxl import Workbook, load_workbook

from common.exceptions import ValidationError
from reference.models import CalcParameterSet, CleaningMethod, GtModel, SmpPrice
from reference.param_defs import default_params

# --- 시드 `[임시값]` (specs/06 §2.1, §3) ---------------------------------------------
# 실존 모델명에 가짜 한계값을 붙이지 않는다 — 실제 사양으로 오해받을 수 있다.
SEED_GT_MODELS: tuple[dict[str, Any], ...] = (
    {
        "name": "가상 모델 A (F급)",
        "rated_gt_mw": 270,
        "rated_st_mw": 135,
        "design_backpressure_kpa": 3.0,
        "backpressure_alarm_kpa": 4.5,
        "backpressure_trip_kpa": 5.5,
        "design_exhaust_temp_c": 610,
        "exhaust_temp_alarm_c": 640,
        "exhaust_temp_trip_c": 660,
        "design_stack_temp_c": 95,
    },
    {
        "name": "가상 모델 B (H급)",
        "rated_gt_mw": 400,
        "rated_st_mw": 200,
        "design_backpressure_kpa": 3.2,
        "backpressure_alarm_kpa": 4.8,
        "backpressure_trip_kpa": 5.8,
        "design_exhaust_temp_c": 640,
        "exhaust_temp_alarm_c": 665,
        "exhaust_temp_trip_c": 685,
        "design_stack_temp_c": 90,
    },
    {
        "name": "가상 모델 C (E급)",
        "rated_gt_mw": 170,
        "rated_st_mw": 85,
        "design_backpressure_kpa": 2.8,
        "backpressure_alarm_kpa": 4.2,
        "backpressure_trip_kpa": 5.2,
        "design_exhaust_temp_c": 560,
        "exhaust_temp_alarm_c": 590,
        "exhaust_temp_trip_c": 610,
        "design_stack_temp_c": 105,
    },
)

SEED_METHODS: tuple[dict[str, Any], ...] = (
    {
        "name": "드라이아이스 세정",
        "cleaning_cost_won": 30_000_000,
        "outage_days": 2.0,
        "recovery_ratio": 0.85,
    },
    {
        "name": "화학 세정",
        "cleaning_cost_won": 80_000_000,
        "outage_days": 4.0,
        "recovery_ratio": 0.95,
    },
    {
        "name": "고압 수세정",
        "cleaning_cost_won": 20_000_000,
        "outage_days": 3.0,
        "recovery_ratio": 0.75,
    },
)

SEED_MANUFACTURER = "가상"
SEED_NOTE = "[임시값] 실무 자료를 받으면 교체한다(specs/06)."

# 시드 SMP 는 '추정' 으로 표시한다 — 확인되지 않은 가격을 공식 가격처럼 보이지 않게(REF-4).
SEED_SMP = {
    "value_won_per_kwh": 150.0,
    "source": "임시값(가상) — 관리자가 KPX 공개 자료로 교체",
    "is_estimate": True,
}


def seed_reference_data() -> dict[str, int]:
    """멱등 시드. 이미 있는 이름은 건드리지 않는다(관리자가 고친 값을 되돌리지 않게)."""
    created = {"gt_models": 0, "methods": 0, "param_sets": 0, "smp": 0}
    for row in SEED_GT_MODELS:
        _, made = GtModel.objects.get_or_create(
            name=row["name"],
            defaults={**row, "manufacturer": SEED_MANUFACTURER, "note": SEED_NOTE},
        )
        created["gt_models"] += made
    for row in SEED_METHODS:
        _, made = CleaningMethod.objects.get_or_create(
            name=row["name"], defaults={**row, "note": SEED_NOTE}
        )
        created["methods"] += made
    if not CalcParameterSet.objects.exists():
        CalcParameterSet.objects.create(
            version_label="p1",
            params=default_params(),
            is_active=True,
            is_seed=True,
            note=SEED_NOTE,
        )
        created["param_sets"] += 1
    if not SmpPrice.objects.exists():
        SmpPrice.objects.create(as_of_date=date.today(), **SEED_SMP)
        created["smp"] += 1
    return created


# --- 파라미터 버전 (specs/06 REF-5) -----------------------------------------------------


def _next_version_label() -> str:
    numbers = [
        int(m.group(1))
        for label in CalcParameterSet.objects.values_list("version_label", flat=True)
        if (m := re.fullmatch(r"p(\d+)", label))
    ]
    return f"p{max(numbers, default=0) + 1}"


@transaction.atomic
def create_param_version(params: dict[str, Any], *, note: str, user) -> CalcParameterSet:
    """새 버전을 만들고 활성화한다. 이전 버전은 남는다."""
    CalcParameterSet.objects.filter(is_active=True).update(is_active=False)
    return CalcParameterSet.objects.create(
        version_label=_next_version_label(),
        params=params,
        is_active=True,
        is_seed=params == default_params(),
        note=note,
        created_by=user,
    )


@transaction.atomic
def activate_param_version(param_set: CalcParameterSet) -> CalcParameterSet:
    CalcParameterSet.objects.filter(is_active=True).exclude(pk=param_set.pk).update(is_active=False)
    param_set.is_active = True
    param_set.save(update_fields=["is_active"])
    return param_set


# --- GT 표 가져오기·내보내기 (specs/06 REF-1, AC-06-6) --------------------------------

GT_COLUMNS = (
    "name",
    "manufacturer",
    "rated_gt_mw",
    "rated_st_mw",
    "design_backpressure_kpa",
    "backpressure_alarm_kpa",
    "backpressure_trip_kpa",
    "design_exhaust_temp_c",
    "exhaust_temp_alarm_c",
    "exhaust_temp_trip_c",
    "design_stack_temp_c",
    "is_placeholder",
    "note",
)
MAX_IMPORT_ROWS = 500
TRUTHY_CELLS = {"true", "1", "y", "yes", "예", "o"}


def gt_template_xlsx() -> bytes:
    """가져오기 양식 — 머리글 + 현재 등록된 모델."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "gt_models"
    sheet.append(list(GT_COLUMNS))
    for model in GtModel.objects.order_by("name"):
        sheet.append([_safe_cell(getattr(model, column)) for column in GT_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _safe_cell(value: Any) -> Any:
    """수식 인젝션 방지(specs/13 §1) — =, +, -, @ 로 시작하는 문자열은 글자로 고정한다."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def _parse_row(raw: dict[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for column in GT_COLUMNS:
        value = raw.get(column)
        if column in {"name", "manufacturer", "note"}:
            row[column] = "" if value is None else str(value).strip()
        elif column == "is_placeholder":
            row[column] = str(value).strip().lower() in TRUTHY_CELLS if value is not None else False
        elif value is None or value == "":
            row[column] = None
        else:
            row[column] = float(value)
    return row


def import_gt_models(file_bytes: bytes) -> dict[str, int]:
    """전부 아니면 전무(AC-06-6). 오류가 하나라도 있으면 행 번호·사유를 알리고 저장하지 않는다."""
    from reference.serializers import GtModelSerializer

    try:
        sheet = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True).active
    except Exception as exc:  # noqa: BLE001 - 손상 파일·암호 파일 모두 같은 안내
        raise ValidationError(message="xlsx 파일을 읽을 수 없습니다.", code="INVALID_FILE") from exc

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValidationError(message="빈 파일입니다.", code="EMPTY_FILE")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in ("name", *GT_COLUMNS[2:11]) if c not in header]
    if missing:
        raise ValidationError(
            message="필수 머리글이 없습니다.", code="MISSING_COLUMNS", details={"missing": missing}
        )
    body = [r for r in rows[1:] if any(cell not in (None, "") for cell in r)]
    if len(body) > MAX_IMPORT_ROWS:
        raise ValidationError(message=f"한 번에 {MAX_IMPORT_ROWS}행까지 가져올 수 있습니다.")

    errors: dict[str, list[str]] = {}
    prepared = []
    for line_no, values in enumerate(body, start=2):
        try:
            data = _parse_row(dict(zip(header, values, strict=False)))
        except (TypeError, ValueError):
            errors[f"{line_no}행"] = ["숫자 칸에 숫자가 아닌 값이 있습니다."]
            continue
        instance = GtModel.objects.filter(name=data["name"]).first()
        serializer = GtModelSerializer(instance, data=data)
        if not serializer.is_valid():
            errors[f"{line_no}행"] = [
                f"{field}: {' '.join(str(m) for m in messages)}"
                for field, messages in serializer.errors.items()
            ]
            continue
        prepared.append(serializer)

    if errors:
        raise ValidationError(
            message="가져올 수 없는 행이 있어 아무것도 저장하지 않았습니다.",
            code="IMPORT_ROWS_INVALID",
            details=errors,
        )

    counts = {"created": 0, "updated": 0}
    try:
        with transaction.atomic():
            for serializer in prepared:
                counts["updated" if serializer.instance else "created"] += 1
                serializer.save()
    except IntegrityError as exc:
        raise ValidationError(message="같은 이름이 파일 안에 두 번 있습니다.") from exc
    return counts
