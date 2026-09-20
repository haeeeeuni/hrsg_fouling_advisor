"""분석 리포트 엑셀 생성 (specs/12 §4).

시트 10종, 열 너비 자동 조정, 헤더 고정, 숫자 서식, 네이티브 FI 차트.
모든 문자열 셀에 CSV 수식 인젝션 방지를 적용한다(AC-18-5).
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reports import formatting as fmt

HEADER_FILL = PatternFill("solid", fgColor="F1F3F5")
HEADER_FONT = Font(bold=True)
MAX_COLUMN_WIDTH = 60


def _write_sheet(
    workbook: Workbook,
    title: str,
    header: list[str],
    rows: list[list[Any]],
    number_formats: dict[int, str] | None = None,
) -> Any:
    sheet = workbook.create_sheet(title)
    sheet.append([fmt.sanitize_cell(value) for value in header])

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([fmt.sanitize_cell(value) for value in row])

    # 숫자 서식 (specs/12 §4.1)
    for index, number_format in (number_formats or {}).items():
        letter = get_column_letter(index)
        for cell in sheet[letter][1:]:
            cell.number_format = number_format

    # 열 너비 자동 조정 + 헤더 고정
    for column_index in range(1, len(header) + 1):
        letter = get_column_letter(column_index)
        width = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
            default=8,
        )
        sheet.column_dimensions[letter].width = min(max(width + 3, 10), MAX_COLUMN_WIDTH)
    sheet.freeze_panes = "A2"
    return sheet


def build_analysis_xlsx(context: dict[str, Any]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    stats = context.get("data_stats") or {}
    trend = context.get("trend") or {}
    benefit = context.get("benefit") or {}
    points = context.get("fouling_points") or []

    # 1. 요약
    _write_sheet(
        workbook,
        "요약",
        ["항목", "값"],
        [
            ["호기", context["unit_name"]],
            ["분석 기간", f"{context['period_start']} ~ {context['period_end']}"],
            ["실행자", context["executed_by"]],
            ["분석 일시", context["executed_at"]],
            ["현재 오염도 지수", context.get("current_fi")],
            ["등급", fmt.grade(context.get("grade"))],
            ["신뢰도", fmt.confidence(context.get("confidence"))],
            ["임계 도달 D-day", fmt.dday(trend.get("eta_days"), trend.get("status"))],
            ["도달 예상일", fmt.ymd(trend.get("eta_date"))],
            ["순편익(원)", benefit.get("net_benefit")],
            ["회수기간(일)", benefit.get("payback_days")],
            [
                "차압 모델",
                f"{(context.get('model_dp') or {}).get('algorithm', '')} "
                f"v{(context.get('model_dp') or {}).get('version', '')}",
            ],
            [
                "스택온도 모델",
                f"{(context.get('model_st') or {}).get('algorithm', '')} "
                f"v{(context.get('model_st') or {}).get('version', '')}",
            ],
        ],
    )

    # 2. 오염도지수 (+ 네이티브 차트)
    fi_sheet = _write_sheet(
        workbook,
        "오염도지수",
        ["일자", "FI", "등급", "표본 수", "신뢰도"],
        [
            [
                fmt.ymd(p["date"]),
                p.get("fi_value"),
                fmt.grade(p.get("grade")),
                p.get("sample_count"),
                fmt.confidence(p.get("confidence")),
            ]
            for p in points
        ],
        {2: "0.0", 4: "#,##0"},
    )
    if len(points) > 1:
        chart = LineChart()
        chart.title = "오염도 지수 시계열"
        chart.y_axis.title = "FI"
        chart.x_axis.title = "일자"
        chart.height, chart.width = 8, 22
        chart.add_data(
            Reference(fi_sheet, min_col=2, min_row=1, max_row=len(points) + 1),
            titles_from_data=True,
        )
        chart.set_categories(Reference(fi_sheet, min_col=1, min_row=2, max_row=len(points) + 1))
        fi_sheet.add_chart(chart, "G2")

    # 3. 잔차상세
    _write_sheet(
        workbook,
        "잔차상세",
        [
            "일자",
            "실측 차압",
            "기대 차압",
            "차압 잔차",
            "실측 스택온도",
            "기대 스택온도",
            "스택온도 잔차",
        ],
        [
            [
                fmt.ymd(p["date"]),
                p.get("measured_dp"),
                p.get("expected_dp"),
                p.get("residual_dp"),
                p.get("measured_st"),
                p.get("expected_st"),
                p.get("residual_st"),
            ]
            for p in points
        ],
        {2: "0.00", 3: "0.00", 4: "0.00", 5: "0.0", 6: "0.0", 7: "0.0"},
    )

    # 4. 군집별분석
    _write_sheet(
        workbook,
        "군집별분석",
        [
            "군집",
            "라벨",
            "표본 수",
            "비중(%)",
            "평균 부하율(%)",
            "평균 외기온(℃)",
            "평균 차압",
            "평균 스택온도",
        ],
        [
            [
                c.get("cluster_key"),
                c.get("label"),
                c.get("sample_count"),
                c.get("share_pct"),
                c.get("avg_load_ratio_pct"),
                c.get("avg_ambient_temp_c"),
                c.get("avg_dp_kpa"),
                c.get("avg_stack_temp_c"),
            ]
            for c in (context.get("clusters") or [])
        ],
        {3: "#,##0", 4: "0.0", 5: "0.0", 6: "0.0", 7: "0.00", 8: "0.0"},
    )

    # 5. 모델정보
    model_rows = []
    for label, model in (("차압", context.get("model_dp")), ("스택온도", context.get("model_st"))):
        if not model:
            continue
        metrics = model.get("metrics") or {}
        model_rows.append(
            [
                label,
                model.get("algorithm"),
                model.get("version"),
                fmt.ymd(model.get("baseline_start")),
                fmt.ymd(model.get("baseline_end")),
                model.get("training_rows"),
                metrics.get("mae"),
                metrics.get("rmse"),
                metrics.get("r2"),
                model.get("residual_std"),
                ", ".join(model.get("feature_list") or []),
                str(model.get("hyperparams") or {}),
            ]
        )
    _write_sheet(
        workbook,
        "모델정보",
        [
            "구분",
            "알고리즘",
            "버전",
            "학습 시작",
            "학습 종료",
            "학습 표본",
            "MAE",
            "RMSE",
            "R²",
            "잔차 σ",
            "피처",
            "하이퍼파라미터",
        ],
        model_rows,
        {6: "#,##0", 7: "0.000", 8: "0.000", 9: "0.000", 10: "0.0000"},
    )

    # 6. 편익분석
    benefit_rows = [
        ["Δ차압(kPa)", benefit.get("delta_dp_kpa")],
        ["Δ스택온도(℃)", benefit.get("delta_stack_c")],
        ["GT 출력 손실(MW)", benefit.get("power_loss_gt_mw")],
        ["ST 출력 손실(MW)", benefit.get("power_loss_st_mw")],
        ["총 출력 손실(MW)", benefit.get("power_loss_total_mw")],
        ["일일 손실 비용(원)", benefit.get("daily_loss_cost")],
        ["일일 연료 손실(원)", benefit.get("daily_fuel_loss")],
        ["세정 비용(원)", benefit.get("cleaning_cost")],
        ["정지 손실(원)", benefit.get("outage_loss")],
        ["총 세정 비용(원)", benefit.get("total_cleaning_cost")],
        ["회수 편익 정밀(원)", benefit.get("gross_benefit")],
        ["회수 편익 간이(원)", benefit.get("gross_benefit_simple")],
        ["순편익(원)", benefit.get("net_benefit")],
        ["회수기간(일)", benefit.get("payback_days")],
        ["ROI(%)", benefit.get("roi_pct")],
        ["권고 세정 일자", fmt.ymd(benefit.get("recommended_cleaning_date"))],
        [None, None],
        ["— 적용 파라미터 —", None],
        *[[k, v] for k, v in sorted((benefit.get("params_snapshot") or {}).items())],
        [None, None],
        ["— 시나리오 —", None],
        *[
            [f"{r['label']} (D+{r.get('offset_days')})", r.get("net_benefit")]
            for r in (benefit.get("scenarios") or [])
        ],
    ]
    _write_sheet(workbook, "편익분석", ["항목", "값"], benefit_rows, {2: "#,##0.00"})

    # 7. 추세예측
    _write_sheet(
        workbook,
        "추세예측",
        ["항목", "값"],
        [
            ["추세 모델", trend.get("model_type")],
            ["상태", fmt.trend_status(trend.get("status"))],
            ["적합 시작", fmt.ymd(trend.get("fit_start"))],
            ["적합 종료", fmt.ymd(trend.get("fit_end"))],
            ["진행률(FI/일)", trend.get("slope_per_day")],
            ["주간 증가량", trend.get("weekly_increase")],
            ["R²", trend.get("r2")],
            ["MAE", trend.get("mae")],
            ["p 값", trend.get("p_value")],
            ["적용 임계치", trend.get("threshold_used")],
            ["D-day", trend.get("eta_days")],
            ["도달 예상일", fmt.ymd(trend.get("eta_date"))],
            ["도달 하한", fmt.ymd(trend.get("eta_lower_date"))],
            ["도달 상한", fmt.ymd(trend.get("eta_upper_date"))],
            ["주의 전환 예상일", fmt.ymd(trend.get("caution_eta_date"))],
            ["경고 전환 예상일", fmt.ymd(trend.get("warning_eta_date"))],
        ],
    )

    # 8. 세정전후비교 (해당 시)
    comparison = context.get("comparison")
    if comparison:
        _write_sheet(
            workbook,
            "세정전후비교",
            ["지표", "세정 전", "세정 후", "개선량", "개선율(%)", "주 지표"],
            [
                [
                    m["label"],
                    m["before"],
                    m["after"],
                    m["delta"],
                    m["delta_pct"],
                    "O" if m["is_primary"] else "",
                ]
                for m in comparison["metrics"]["rows"]
            ],
        )

    # 9. 데이터품질
    _write_sheet(
        workbook,
        "데이터품질",
        ["항목", "값"],
        [
            ["총 포인트", stats.get("row_total")],
            ["유효 포인트", stats.get("row_valid")],
            ["유효 비율", stats.get("valid_ratio")],
            ["유효 세그먼트 수", stats.get("segment_count")],
            ["청정 기준 포인트", stats.get("baseline_points")],
            ["청정 기준 결정 경로", stats.get("baseline_source")],
            ["도메인 밖 비율", stats.get("out_of_domain_ratio")],
            [None, None],
            ["— 제외 사유별 —", None],
            *[[k, v] for k, v in (stats.get("excluded_by_reason") or {}).items()],
        ],
        {2: "#,##0.0000"},
    )

    # 10. 설정값스냅샷
    _write_sheet(
        workbook,
        "설정값스냅샷",
        ["키", "값"],
        [[k, str(v)] for k, v in sorted((context.get("settings_snapshot") or {}).items())],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
