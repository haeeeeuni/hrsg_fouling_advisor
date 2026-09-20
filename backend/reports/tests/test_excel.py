"""엑셀 리포트 검증 (specs/12 §4, §7). DB 불필요."""

import io

import pytest
from openpyxl import load_workbook

from reports import formatting as fmt
from reports.excel.builder import build_analysis_xlsx

EXPECTED_SHEETS = [
    "요약",
    "오염도지수",
    "잔차상세",
    "군집별분석",
    "모델정보",
    "편익분석",
    "추세예측",
    "데이터품질",
    "설정값스냅샷",
]


@pytest.fixture
def workbook(context):
    return load_workbook(io.BytesIO(build_analysis_xlsx(context)))


def test_ac_12_3_file_opens_as_valid_xlsx(context):
    """AC-12-3: 엑셀 파일이 정상 열린다(openpyxl 로 재해석 가능)."""
    content = build_analysis_xlsx(context)

    assert content[:2] == b"PK"  # xlsx 는 zip 컨테이너
    assert load_workbook(io.BytesIO(content)).sheetnames


def test_required_sheets_exist(workbook):
    for sheet in EXPECTED_SHEETS:
        assert sheet in workbook.sheetnames


def test_comparison_sheet_appears_only_when_present(context, comparison_context):
    plain = load_workbook(io.BytesIO(build_analysis_xlsx(context)))
    with_comparison = load_workbook(io.BytesIO(build_analysis_xlsx(comparison_context)))

    assert "세정전후비교" not in plain.sheetnames
    assert "세정전후비교" in with_comparison.sheetnames


def test_korean_headers_and_values_survive(workbook):
    summary = workbook["요약"]
    values = [row[0] for row in summary.iter_rows(min_col=1, max_col=1, values_only=True)]

    assert "호기" in values
    assert "현재 오염도 지수" in values
    assert summary["A1"].value == "항목"


def test_native_chart_is_embedded(workbook):
    """specs/12 §4.2 — 엑셀 네이티브 차트를 최소 1개 포함한다."""
    assert len(workbook["오염도지수"]._charts) >= 1


def test_header_is_frozen(workbook):
    for sheet in EXPECTED_SHEETS:
        assert workbook[sheet].freeze_panes == "A2"


def test_number_formats_are_applied(workbook):
    sheet = workbook["오염도지수"]

    assert sheet.cell(row=2, column=2).number_format == "0.0"
    assert sheet.cell(row=2, column=4).number_format == "#,##0"


def test_column_widths_are_adjusted(workbook):
    widths = workbook["요약"].column_dimensions

    assert widths["A"].width >= 10
    assert widths["A"].width <= 60


def test_fouling_points_are_written(workbook, points):
    assert workbook["오염도지수"].max_row == len(points) + 1


def test_settings_snapshot_sheet_has_every_key(workbook, context):
    sheet = workbook["설정값스냅샷"]
    keys = [row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)]

    assert set(keys) == set(context["settings_snapshot"])


# --- AC-18-5 : CSV 수식 인젝션 ---


@pytest.mark.parametrize("cell", ["=1+1", "+CMD", "-SUM(A1)", "@import"])
def test_ac_18_5_formula_injection_is_neutralized(cell):
    assert fmt.sanitize_cell(cell).startswith("'")


def test_numbers_are_not_quoted():
    """음수를 문자열로 바꾸면 계산이 깨진다."""
    assert fmt.sanitize_cell(-1500) == -1500
    assert fmt.sanitize_cell(3.14) == 3.14


def test_ordinary_text_is_untouched():
    assert fmt.sanitize_cell("정상") == "정상"


def test_injection_is_applied_in_generated_file(context):
    poisoned = {**context, "settings_snapshot": {"=cmd|calc": "=1+1"}}

    workbook = load_workbook(io.BytesIO(build_analysis_xlsx(poisoned)))
    sheet = workbook["설정값스냅샷"]

    assert sheet.cell(row=2, column=1).value.startswith("'")
    assert sheet.cell(row=2, column=2).value.startswith("'")
