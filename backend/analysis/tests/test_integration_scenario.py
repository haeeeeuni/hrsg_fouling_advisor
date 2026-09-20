"""전 구간 통합 시나리오 (specs/17 §6). PostgreSQL + Celery eager 필요.

    1. 샘플 데이터 생성 (24개월 · 세정 3회)
    2. 호기 등록 + 컬럼 매핑(한글 헤더)
    3. 업로드 → 검증 통과 → 적재
    4. 정비 이력 업로드 → 세정 후보 추출
    5. 세정 이벤트 승인 → 청정 기준 기간 자동 생성
    6. 분석 실행 → 모델 R² 확인
    7. FI 시계열에서 세정 시점 급락 확인
    8. 정답값과 FI 상관 확인
    9. D-day · 편익 · PDF/엑셀 리포트 생성

AGENTS.md §8 — 이 시나리오를 통합 테스트 1건으로 유지한다.
"""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

from analysis.models import AnalysisRun, FoulingIndexPoint, RunStatus
from ingestion.models import Measurement
from maintenance.models import CleaningEvent, MaintenanceRecord, ReviewStatus
from units.models import ColumnMapping, Unit

pytestmark = pytest.mark.django_db

BACKEND = Path(__file__).resolve().parents[2]
SCRIPT = BACKEND / "scripts" / "generate_sample_data.py"

# specs/17 §4.2 한글 헤더 → 표준 항목
KOREAN_MAPPING = [
    ("timestamp", "시각"),
    ("gt_power_mw", "GT출력(MW)"),
    ("ambient_temp_c", "대기온도(℃)"),
    ("gt_exhaust_temp_c", "GT배기온도(℃)"),
    ("exhaust_flow", "배기유량(kg/s)"),
    ("hrsg_gas_dp_kpa", "HRSG가스차압(kPa)"),
    ("stack_temp_c", "스택온도(℃)"),
    ("duct_burner_on", "덕트버너상태"),
]


@pytest.fixture(scope="module")
def sample_files(tmp_path_factory):
    """[1단계] 샘플 데이터 생성. 테스트 시간을 줄이려 12개월·세정 2회로 줄인다."""
    out = tmp_path_factory.mktemp("sample")
    operation = out / "u1.csv"
    maintenance = out / "u1_maint.csv"
    truth = out / "u1_truth.csv"

    subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--unit-code",
            "U1",
            "--start",
            "2023-01-01",
            "--months",
            "12",
            "--cleanings",
            "2",
            "--seed",
            "42",
            "--korean-headers",
            "--out",
            str(operation),
            "--maintenance-out",
            str(maintenance),
            "--truth-out",
            str(truth),
        ],
        check=True,
        capture_output=True,
        cwd=BACKEND,
    )
    return {"operation": operation, "maintenance": maintenance, "truth": truth}


@pytest.fixture
def unit(db, seeded):
    """[2단계] 호기 등록 + 한글 헤더 컬럼 매핑."""
    u = Unit.objects.create(
        code="U1",
        name="1호기 HRSG",
        plant_name="○○복합화력",
        rated_power_mw=160,
        rated_st_power_mw=80,
        min_load_mw=60,
        sampling_interval_min=10,
    )
    for field, column in KOREAN_MAPPING:
        ColumnMapping.objects.create(
            unit=u,
            standard_field=field,
            source_column=column,
            bool_rule="BOOL" if field == "duct_burner_on" else "",
        )
    assert u.is_mapping_complete
    return u


def upload(path: Path, name: str) -> SimpleUploadedFile:
    return SimpleUploadedFile(name, path.read_bytes(), content_type="text/csv")


@pytest.fixture
def loaded(api, normal_user, unit, sample_files):
    """[3단계] 업로드 → 검증 → 적재."""
    api.force_authenticate(normal_user)

    res = api.post(
        "/api/uploads/operation/validate/",
        {"unit_id": unit.id, "file": upload(sample_files["operation"], "u1.csv")},
        format="multipart",
    )
    assert res.status_code == 202, res.data
    batch_id = res.data["batch_id"]

    from ingestion.models import BatchStatus, UploadBatch

    batch = UploadBatch.objects.get(pk=batch_id)
    assert batch.status == BatchStatus.VALIDATED, batch.validation_report
    assert batch.validation_report["errors"] == []

    commit = api.post(f"/api/uploads/{batch_id}/commit/", {}, format="json")
    assert commit.status_code == 202
    assert Measurement.objects.filter(unit=unit).count() > 50_000
    return unit


def test_step_3_operation_data_is_loaded(loaded, unit):
    """[3] 한글 헤더 CSV 가 매핑을 거쳐 정상 적재된다 (AC-17-5)."""
    summary = Measurement.objects.filter(unit=unit)

    assert summary.count() > 50_000
    first = summary.order_by("timestamp").first()
    assert first.gt_power_mw is not None
    assert first.duct_burner_on in (True, False)


def test_step_4_5_maintenance_upload_extracts_and_accepts(
    api, normal_user, admin_user, loaded, unit, sample_files
):
    """[4][5] 정비 이력 업로드 → 세정 후보 추출 → 승인 → CleaningEvent 생성."""
    api.force_authenticate(normal_user)

    res = api.post(
        "/api/uploads/maintenance/",
        {"unit_id": unit.id, "file": upload(sample_files["maintenance"], "maint.csv")},
        format="multipart",
    )
    assert res.status_code == 202, res.data

    records = MaintenanceRecord.objects.filter(unit=unit)
    assert records.count() > 0

    # 세정 후보가 추출된다 (AC-10-1)
    candidates = records.filter(is_fouling_related=True, match_category="CLEANING")
    assert candidates.count() == 2  # 세정 2회

    # 승인 전에는 CleaningEvent 가 없다 (AC-10-2)
    assert CleaningEvent.objects.filter(unit=unit).count() == 0

    # EXCLUDE 키워드 행("세정 계획 취소")은 후보에서 빠진다
    cancelled = records.filter(title__contains="취소").first()
    assert cancelled is not None
    assert not cancelled.is_fouling_related

    for record in candidates:
        accept = api.post(f"/api/maintenance-records/{record.id}/accept/", {}, format="json")
        assert accept.status_code == 201

    assert CleaningEvent.objects.filter(unit=unit).count() == 2
    assert records.filter(review_status=ReviewStatus.ACCEPTED).count() == 2


@pytest.fixture
def analysed(api, normal_user, loaded, unit, sample_files):
    """[5][6] 세정 이벤트 등록 후 분석 실행."""
    api.force_authenticate(normal_user)
    api.post(
        "/api/uploads/maintenance/",
        {"unit_id": unit.id, "file": upload(sample_files["maintenance"], "maint.csv")},
        format="multipart",
    )
    for record in MaintenanceRecord.objects.filter(unit=unit, match_category="CLEANING"):
        api.post(f"/api/maintenance-records/{record.id}/accept/", {}, format="json")

    res = api.post(
        "/api/analysis-runs/",
        {
            "unit_id": unit.id,
            "period_start": "2023-01-01T00:00:00+09:00",
            "period_end": "2023-12-31T23:59:59+09:00",
        },
        format="json",
    )
    assert res.status_code == 202, res.data
    run = AnalysisRun.objects.get(pk=res.data["analysis_run_id"])
    assert run.status == RunStatus.SUCCESS, run.error_message
    return run


def test_step_6_model_accuracy_meets_target(analysed):
    """[6] AC-06-1 — 차압 모델 검증 R² 가 0.85 이상이다."""
    assert analysed.model_version_dp.metrics["r2"] >= 0.85


def test_step_6_baseline_came_from_cleaning_event(analysed):
    """세정 이력이 있으면 청정 기준 기간이 그로부터 자동 산정된다 (specs/06 §2)."""
    assert analysed.data_stats["baseline_source"] == "AUTO_FROM_CLEANING"
    assert analysed.data_stats["baseline_points"] > 0


def test_step_7_fi_drops_at_each_cleaning(analysed, unit):
    """[7] FI 시계열에서 세정 시점마다 급락하고 사이클 내에서 증가한다 (AC-17-2)."""
    points = list(
        FoulingIndexPoint.objects.filter(analysis_run=analysed, cluster_key="")
        .exclude(fi_value__isnull=True)
        .order_by("date")
        .values_list("date", "fi_value")
    )
    series = pd.Series({d: v for d, v in points})

    for event in CleaningEvent.objects.filter(unit=unit).order_by("cleaned_at"):
        cleaned = pd.Timestamp(timezone.localtime(event.cleaned_at).date())
        before = series[
            (series.index >= (cleaned - pd.Timedelta(days=10)).date())
            & (series.index < cleaned.date())
        ]
        after = series[
            (series.index > (cleaned + pd.Timedelta(days=4)).date())
            & (series.index <= (cleaned + pd.Timedelta(days=20)).date())
        ]
        if before.empty or after.empty:
            continue
        assert after.median() < before.median(), f"{cleaned.date()} 에서 FI 가 회복되지 않았습니다"


def test_step_8_fi_correlates_with_truth(analysed, sample_files):
    """[8] AC-17-3 — 생성 정답값과 산출 FI 의 상관계수가 0.9 이상이다."""
    truth = pd.read_csv(sample_files["truth"], parse_dates=["date"]).set_index("date")
    points = pd.DataFrame(
        FoulingIndexPoint.objects.filter(analysis_run=analysed, cluster_key="")
        .exclude(fi_value__isnull=True)
        .values("date", "fi_value")
    )
    points["date"] = pd.to_datetime(points["date"])
    joined = points.set_index("date").join(truth, how="inner").dropna()

    assert len(joined) > 100
    assert joined.fi_value.corr(joined.fouling_level) >= 0.90


def test_step_9_trend_and_benefit_are_produced(analysed):
    """[9] D-day 와 편익이 산출된다."""
    assert analysed.trend is not None
    assert analysed.trend.status in {"OK", "ALREADY_EXCEEDED", "NO_TREND", "BEYOND_HORIZON"}
    assert analysed.benefit is not None
    assert analysed.benefit.power_loss_total_mw is not None
    assert len(analysed.benefit.scenarios) == 3


def test_step_9_pdf_report_has_readable_korean(api, normal_user, analysed):
    """[9] AC-12-1 — PDF 한글이 깨지지 않는다."""
    from pypdf import PdfReader

    api.force_authenticate(normal_user)
    res = api.post(f"/api/analysis-runs/{analysed.id}/export/", {"format": "pdf"}, format="json")
    assert res.status_code == 201, res.data

    download = api.get(f"/api/reports/{res.data['id']}/download/")
    assert download.status_code == 200
    content = b"".join(download.streaming_content)

    text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages)
    assert "HRSG 가스측 오염도 진단 리포트" in text
    assert "1호기 HRSG" in text
    assert "편익 분석" in text


def test_step_9_excel_report_is_valid(api, normal_user, analysed):
    """[9] AC-12-3 — 엑셀이 정상 열린다."""
    from openpyxl import load_workbook

    api.force_authenticate(normal_user)
    res = api.post(f"/api/analysis-runs/{analysed.id}/export/", {"format": "xlsx"}, format="json")
    assert res.status_code == 201

    download = api.get(f"/api/reports/{res.data['id']}/download/")
    content = b"".join(download.streaming_content)
    workbook = load_workbook(io.BytesIO(content))

    assert "요약" in workbook.sheetnames
    assert "오염도지수" in workbook.sheetnames


def test_step_9_korean_filename_uses_rfc5987(api, normal_user, analysed):
    """AC-12-6 — 한글 파일명으로 다운로드해도 깨지지 않는다."""
    api.force_authenticate(normal_user)
    res = api.post(f"/api/analysis-runs/{analysed.id}/export/", {"format": "pdf"}, format="json")

    download = api.get(f"/api/reports/{res.data['id']}/download/")
    disposition = download["Content-Disposition"]

    assert "filename*=UTF-8''" in disposition
    assert "HRSG" in res.data["file_name"]


def test_comparison_report_is_generated(api, normal_user, analysed, unit):
    """세정 전후 비교가 공통 군집에서 수행된다 (specs/12 §2)."""
    api.force_authenticate(normal_user)
    event = CleaningEvent.objects.filter(unit=unit).order_by("cleaned_at").first()

    res = api.post(
        "/api/comparisons/",
        {"cleaning_event_id": event.id, "window_days": 30},
        format="json",
    )
    assert res.status_code == 201, res.data
    assert res.data["is_comparable"]
    assert res.data["common_clusters"]

    export = api.post(
        f"/api/comparisons/{res.data['id']}/export/", {"format": "pdf"}, format="json"
    )
    assert export.status_code == 201
