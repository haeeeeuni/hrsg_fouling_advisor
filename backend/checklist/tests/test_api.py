"""체크리스트 API (specs/07 AC-07-1~6, specs/10 §5·§6)."""

import io

import pytest
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from accounts.models import User
from checklist.models import ChecklistTemplateItem, DataRequest
from checklist.seed import SEED_ITEMS
from checklist.services import TEMPLATE_COLUMNS
from common.models import AuditLog

pytestmark = pytest.mark.django_db

REQUESTS = "/api/checklist/requests/"
ITEMS = "/api/admin/checklist-items/"


@pytest.fixture
def as_user(api, normal_user, seeded):
    api.force_authenticate(normal_user)
    return api


@pytest.fixture
def created(as_user):
    return as_user.post(REQUESTS, {"title": "10월 A 사이트 사전 평가"}, format="json").json()


def set_state(api, request_id, item_id, state):
    return api.patch(f"{REQUESTS}{request_id}/items/{item_id}/", {"state": state}, format="json")


# --- 시드 (AC-07-6) ---


def test_seed_has_pinch_and_schedule_items_once(seeded):
    call_command("seed_defaults", verbosity=0)

    assert ChecklistTemplateItem.objects.count() == len(SEED_ITEMS)
    pinch = set(
        ChecklistTemplateItem.objects.filter(category="PINCH").values_list("name_ko", flat=True)
    )
    assert pinch == {"압력단별 드럼 압력", "절탄기 출구 급수 온도", "증발기 출구 가스 온도"}
    assert ChecklistTemplateItem.objects.filter(category="SCHEDULE").count() >= 4
    assert all(i.is_placeholder for i in ChecklistTemplateItem.objects.all())


# --- 요청 건 ---


def test_new_request_copies_active_template(created):
    """AC-07-1"""
    assert len(created["items"]) == len(SEED_ITEMS)
    assert created["progress"]["percent"] == 0
    assert created["status"] == "IN_PROGRESS"
    assert created["new_template_items"] == 0


def test_inactive_template_items_are_not_copied(as_user):
    ChecklistTemplateItem.objects.filter(category="DATA_SCOPE").update(is_active=False)

    data = as_user.post(REQUESTS, {"title": "x"}, format="json").json()

    assert all(i["category"] != "DATA_SCOPE" for i in data["items"])


def test_title_is_required(as_user):
    assert as_user.post(REQUESTS, {"title": "  "}, format="json").status_code == 400


def test_receiving_all_required_items_completes_the_request(as_user, created):
    """AC-07-2 — 필수를 다 받으면 완료, '해당 없음' 은 분모에서 빠진다."""
    rid = created["id"]
    required = [i for i in created["items"] if i["is_required"]]
    for item in required[:-1]:
        set_state(as_user, rid, item["id"], "RECEIVED")
    last = set_state(as_user, rid, required[-1]["id"], "NOT_APPLICABLE").json()

    assert last["request"]["status"] == "DONE"
    assert last["request"]["progress"]["complete"] is True

    reopened = set_state(as_user, rid, required[0]["id"], "PENDING").json()
    assert reopened["request"]["status"] == "IN_PROGRESS"


def test_received_at_is_recorded_and_cleared(as_user, created):
    item = created["items"][0]

    received = set_state(as_user, created["id"], item["id"], "RECEIVED").json()["item"]
    assert received["received_at"]

    pending = set_state(as_user, created["id"], item["id"], "PENDING").json()["item"]
    assert pending["received_at"] is None


def test_item_memo(as_user, created):
    item = created["items"][0]

    res = as_user.patch(
        f"{REQUESTS}{created['id']}/items/{item['id']}/",
        {"memo": "10/15 메일로 받기로"},
        format="json",
    )

    assert res.json()["item"]["memo"] == "10/15 메일로 받기로"


def test_invalid_state_is_rejected(as_user, created):
    assert set_state(as_user, created["id"], created["items"][0]["id"], "DONE").status_code == 400


def test_template_changes_do_not_alter_existing_requests(as_user, created, admin_user, api):
    """AC-07-3 — 그리고 새 항목은 사용자가 고를 때만 덧붙는다."""
    template = ChecklistTemplateItem.objects.get(name_ko="GT 출력")
    template.name_ko = "GT 출력(수정됨)"
    template.save()
    ChecklistTemplateItem.objects.create(
        category="GT", name_ko="새 항목", name_en="New item", is_required=True, order=999
    )

    detail = as_user.get(f"{REQUESTS}{created['id']}/").json()
    assert "GT 출력" in [i["name_ko"] for i in detail["items"]]
    assert "GT 출력(수정됨)" not in [i["name_ko"] for i in detail["items"]]
    assert detail["new_template_items"] == 1

    synced = as_user.post(f"{REQUESTS}{created['id']}/sync-template/").json()
    assert synced["added"] == 1
    assert synced["request"]["new_template_items"] == 0
    assert len(synced["request"]["items"]) == len(created["items"]) + 1


def test_other_users_request_is_404(api, created, user_password):
    """AC-07-5"""
    other = User.objects.create_user(
        username="other.user",
        full_name="다른 사람",
        organization="x",
        password=user_password,
        approval_status="APPROVED",
    )
    api.force_authenticate(other)
    rid = created["id"]
    item_id = created["items"][0]["id"]

    assert api.get(f"{REQUESTS}{rid}/").status_code == 404
    assert set_state(api, rid, item_id, "RECEIVED").status_code == 404
    assert api.get(f"{REQUESTS}{rid}/email-text/").status_code == 404
    assert api.delete(f"{REQUESTS}{rid}/").status_code == 404
    assert api.get(REQUESTS).json()["count"] == 0


def test_archive_hides_from_default_list_and_unarchive_restores_status(as_user, created):
    rid = created["id"]
    assert (
        as_user.patch(f"{REQUESTS}{rid}/", {"status": "ARCHIVED"}, format="json").status_code == 200
    )
    assert as_user.get(REQUESTS).json()["count"] == 0
    assert as_user.get(REQUESTS, {"status": "ARCHIVED"}).json()["count"] == 1

    restored = as_user.patch(f"{REQUESTS}{rid}/", {"status": "IN_PROGRESS"}, format="json").json()
    assert restored["status"] == "IN_PROGRESS"


def test_email_text_languages_and_scopes(as_user, created):
    """AC-07-4"""
    rid = created["id"]
    gt_output = next(i for i in created["items"] if i["name_ko"] == "GT 출력")
    set_state(as_user, rid, gt_output["id"], "RECEIVED")

    pending_ko = as_user.get(f"{REQUESTS}{rid}/email-text/", {"lang": "ko"}).json()["text"]
    assert "[가스터빈]" in pending_ko
    assert "- GT 출력 (MW)" not in pending_ko
    assert "안녕하십니까" in pending_ko

    all_en = as_user.get(f"{REQUESTS}{rid}/email-text/", {"lang": "en", "scope": "all"}).json()[
        "text"
    ]
    assert "[Gas Turbine]" in all_en
    assert "- GT power output (MW) [Required]" in all_en

    assert as_user.get(f"{REQUESTS}{rid}/email-text/", {"lang": "jp"}).status_code == 400


def test_user_with_requests_cannot_be_hard_deleted(api, admin_user, normal_user, created):
    api.force_authenticate(admin_user)

    res = api.delete(f"/api/admin/users/{normal_user.id}/?hard=true")

    assert res.status_code == 409
    assert DataRequest.objects.filter(owner=normal_user).exists()


# --- 관리자 템플릿 (CHK-6) ---


@pytest.fixture
def as_admin(api, admin_user, seeded):
    api.force_authenticate(admin_user)
    return api


def test_admin_creates_item_at_the_end(as_admin):
    res = as_admin.post(
        ITEMS, {"category": "GT", "name_ko": "추가 항목", "name_en": "Extra"}, format="json"
    )

    assert res.status_code == 201
    assert res.data["order"] > max(
        ChecklistTemplateItem.objects.exclude(pk=res.data["id"]).values_list("order", flat=True)
    )


def test_duplicate_name_in_category_is_rejected(as_admin):
    res = as_admin.post(
        ITEMS, {"category": "GT", "name_ko": "GT 출력", "name_en": "dup"}, format="json"
    )

    assert res.status_code == 400
    assert "같은 분류" in res.json()["error"]["details"]["name_ko"][0]


def test_delete_deactivates_item(as_admin):
    item = ChecklistTemplateItem.objects.first()

    as_admin.delete(f"{ITEMS}{item.id}/")

    item.refresh_from_db()
    assert item.is_active is False
    assert not AuditLog.objects.filter(
        target_type="ChecklistTemplateItem", action="DELETE"
    ).exists()


def test_reorder(as_admin):
    ids = list(ChecklistTemplateItem.objects.order_by("order").values_list("id", flat=True))[:3]

    res = as_admin.post(f"{ITEMS}reorder/", {"ids": list(reversed(ids))}, format="json")

    assert res.status_code == 200
    orders = dict(ChecklistTemplateItem.objects.filter(pk__in=ids).values_list("id", "order"))
    assert orders[ids[2]] < orders[ids[1]] < orders[ids[0]]


def xlsx(rows, columns=TEMPLATE_COLUMNS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(columns))
    for row in rows:
        sheet.append([row.get(c) for c in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    buffer.name = "items.xlsx"
    return buffer


def test_import_upserts_and_marks_as_real(as_admin):
    rows = [
        {
            "category": "GT",
            "name_ko": "GT 출력",
            "name_en": "GT power output (updated)",
            "unit": "MW",
            "is_required": "예",
        },
        {
            "category": "FUEL",
            "name_ko": "연료 가격",
            "name_en": "Fuel price",
            "is_required": "아니오",
        },
    ]

    res = as_admin.post(f"{ITEMS}import/", {"file": xlsx(rows)}, format="multipart")

    assert res.status_code == 200, res.content
    assert res.data == {"created": 1, "updated": 1}
    updated = ChecklistTemplateItem.objects.get(category="GT", name_ko="GT 출력")
    assert updated.name_en == "GT power output (updated)"
    assert updated.is_placeholder is False
    assert ChecklistTemplateItem.objects.get(name_ko="연료 가격").is_required is False


def test_import_is_all_or_nothing(as_admin):
    rows = [
        {"category": "GT", "name_ko": "정상 행", "name_en": "ok"},
        {"category": "NOPE", "name_ko": "오류 행", "name_en": "bad"},
    ]

    res = as_admin.post(f"{ITEMS}import/", {"file": xlsx(rows)}, format="multipart")

    assert res.status_code == 400
    assert "3행" in res.json()["error"]["details"]
    assert not ChecklistTemplateItem.objects.filter(name_ko="정상 행").exists()


def test_export_round_trips(as_admin):
    sheet = load_workbook(io.BytesIO(as_admin.get(f"{ITEMS}export/").content)).active
    rows = list(sheet.iter_rows(values_only=True))

    assert list(rows[0]) == list(TEMPLATE_COLUMNS)
    assert len(rows) == 1 + ChecklistTemplateItem.objects.count()
