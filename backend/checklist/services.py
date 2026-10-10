"""체크리스트 도메인 로직 — 요청 건 생성·항목 상태·템플릿 동기화·시드·가져오기 (specs/07)."""

from __future__ import annotations

import io
from typing import Any

from django.db import IntegrityError, transaction
from django.db.models import Max
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from checklist import rules
from checklist.models import (
    CATEGORY_EN,
    CATEGORY_ORDER,
    Category,
    ChecklistTemplateItem,
    DataRequest,
    DataRequestItem,
    ItemState,
    RequestStatus,
)
from checklist.seed import SEED_ITEMS
from common.exceptions import ValidationError
from common.settings_resolver import get_setting

COPIED_FIELDS = (
    "category",
    "name_ko",
    "name_en",
    "unit",
    "is_required",
    "why_needed_ko",
    "why_needed_en",
    "is_calculator_input",
    "order",
)


def active_template():
    return ChecklistTemplateItem.objects.filter(is_active=True).order_by("order", "id")


def _copy(template: ChecklistTemplateItem, request: DataRequest) -> DataRequestItem:
    return DataRequestItem(
        request=request,
        template_item=template,
        **{field: getattr(template, field) for field in COPIED_FIELDS},
    )


@transaction.atomic
def create_request(*, owner, title: str, memo: str = "", due_date=None) -> DataRequest:
    """요청 건을 만들고 활성 템플릿 항목을 복사한다(CHK-3, AC-07-1)."""
    request = DataRequest.objects.create(
        owner=owner,
        title=title,
        memo=memo,
        due_date=due_date,
        template_snapshot_at=timezone.now(),
    )
    DataRequestItem.objects.bulk_create([_copy(t, request) for t in active_template()])
    return request


def item_dicts(request: DataRequest) -> list[dict[str, Any]]:
    return list(
        request.items.values(
            "category",
            "name_ko",
            "name_en",
            "unit",
            "is_required",
            "why_needed_ko",
            "why_needed_en",
            "state",
        )
    )


def refresh_status(request: DataRequest) -> None:
    """필수 항목을 다 받으면 자동으로 완료, 다시 미수신이 생기면 진행 중으로.

    보관은 사용자가 정한다.
    """
    if request.status == RequestStatus.ARCHIVED:
        return
    complete = rules.progress(item_dicts(request))["complete"]
    status = RequestStatus.DONE if complete else RequestStatus.IN_PROGRESS
    if request.status != status:
        request.status = status
        request.save(update_fields=["status", "updated_at"])


@transaction.atomic
def update_item(item: DataRequestItem, *, state: str | None, memo: str | None) -> DataRequestItem:
    fields = []
    if state is not None and state != item.state:
        item.state = state
        item.received_at = timezone.now() if state == ItemState.RECEIVED else None
        fields += ["state", "received_at"]
    if memo is not None and memo != item.memo:
        item.memo = memo
        fields.append("memo")
    if fields:
        item.save(update_fields=fields)
        DataRequest.objects.filter(pk=item.request_id).update(updated_at=timezone.now())
        refresh_status(item.request)
    return item


def new_template_items(request: DataRequest):
    """요청 건을 만든 뒤 템플릿에 새로 생긴 활성 항목."""
    existing = request.items.exclude(template_item=None).values_list("template_item_id", flat=True)
    return active_template().exclude(pk__in=list(existing))


@transaction.atomic
def sync_template(request: DataRequest) -> int:
    """새 항목만 덧붙인다. 기존 항목·체크는 건드리지 않는다(사용자가 선택, specs/07 CHK-3)."""
    created = DataRequestItem.objects.bulk_create(
        [_copy(t, request) for t in new_template_items(request)]
    )
    if created:
        refresh_status(request)
    return len(created)


def email_text(request: DataRequest, *, lang: str, scope: str) -> str:
    labels = (
        {code: CATEGORY_EN[code] for code in Category.values}
        if lang == rules.LANG_EN
        else dict(Category.choices)
    )
    return rules.email_text(
        item_dicts(request),
        lang=lang,
        scope=scope,
        header=get_setting(f"chk_email_header_{lang}"),
        footer=get_setting(f"chk_email_footer_{lang}"),
        category_labels=labels,
        category_order=CATEGORY_ORDER,
    )


# --- 시드 `[임시값]` ------------------------------------------------------------------


def seed_template() -> int:
    """멱등. 같은 (분류, 이름)이 있으면 건드리지 않는다(관리자가 고친 값을 되돌리지 않게)."""
    created = 0
    for order, row in enumerate(SEED_ITEMS, start=1):
        category, name_ko, name_en, unit, required, why_ko, why_en, calc_input, source = row
        _, made = ChecklistTemplateItem.objects.get_or_create(
            category=category,
            name_ko=name_ko,
            defaults={
                "name_en": name_en,
                "unit": unit,
                "is_required": required,
                "why_needed_ko": why_ko,
                "why_needed_en": why_en,
                "is_calculator_input": calc_input,
                "source": source,
                "order": order * 10,
            },
        )
        created += made
    return created


def next_order() -> int:
    return (ChecklistTemplateItem.objects.aggregate(m=Max("order"))["m"] or 0) + 10


# --- 템플릿 xlsx 가져오기·내보내기 (specs/07 CHK-6) ------------------------------------

TEMPLATE_COLUMNS = (
    "category",
    "name_ko",
    "name_en",
    "unit",
    "is_required",
    "why_needed_ko",
    "why_needed_en",
    "is_calculator_input",
    "source",
    "order",
    "is_active",
)
BOOL_COLUMNS = {"is_required", "is_calculator_input", "is_active"}
TRUTHY_CELLS = {"true", "1", "y", "yes", "예", "o"}
MAX_IMPORT_ROWS = 1000


def _safe_cell(value: Any) -> Any:
    """수식 인젝션 방지(specs/13 §1)."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def export_template_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "checklist_items"
    sheet.append(list(TEMPLATE_COLUMNS))
    for item in ChecklistTemplateItem.objects.order_by("order", "id"):
        sheet.append([_safe_cell(getattr(item, column)) for column in TEMPLATE_COLUMNS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parse_row(raw: dict[str, Any]) -> dict[str, Any]:
    row: dict[str, Any] = {}
    for column in TEMPLATE_COLUMNS:
        value = raw.get(column)
        if column in BOOL_COLUMNS:
            default = column != "is_calculator_input"
            row[column] = (
                default if value in (None, "") else str(value).strip().lower() in TRUTHY_CELLS
            )
        elif column == "order":
            row[column] = None if value in (None, "") else int(float(value))
        else:
            row[column] = "" if value is None else str(value).strip()
    if not row["source"]:
        row["source"] = "FORM"
    return row


def import_template_xlsx(file_bytes: bytes) -> dict[str, int]:
    """전부 아니면 전무. (분류, 한국어 이름)이 같으면 수정, 없으면 추가한다."""
    from checklist.serializers import TemplateItemSerializer

    try:
        sheet = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True).active
    except Exception as exc:  # noqa: BLE001 - 손상 파일·암호 파일 모두 같은 안내
        raise ValidationError(message="xlsx 파일을 읽을 수 없습니다.", code="INVALID_FILE") from exc

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValidationError(message="빈 파일입니다.", code="EMPTY_FILE")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c in ("category", "name_ko", "name_en") if c not in header]
    if missing:
        raise ValidationError(
            message="필수 머리글이 없습니다.", code="MISSING_COLUMNS", details={"missing": missing}
        )
    body = [r for r in rows[1:] if any(cell not in (None, "") for cell in r)]
    if len(body) > MAX_IMPORT_ROWS:
        raise ValidationError(message=f"한 번에 {MAX_IMPORT_ROWS}행까지 가져올 수 있습니다.")

    errors: dict[str, list[str]] = {}
    prepared = []
    order = next_order()
    for line_no, values in enumerate(body, start=2):
        try:
            data = _parse_row(dict(zip(header, values, strict=False)))
        except (TypeError, ValueError):
            errors[f"{line_no}행"] = ["order 칸은 숫자여야 합니다."]
            continue
        instance = ChecklistTemplateItem.objects.filter(
            category=data["category"], name_ko=data["name_ko"]
        ).first()
        if data["order"] is None:
            data["order"] = instance.order if instance else order
            order += 10
        serializer = TemplateItemSerializer(instance, data={**data, "is_placeholder": False})
        if not serializer.is_valid():
            errors[f"{line_no}행"] = [
                f"{field}: {' '.join(str(m) for m in messages)}"
                for field, messages in serializer.errors.items()
            ]
            continue
        prepared.append(serializer)

    if errors:
        raise ValidationError(
            message="가져올 수 없는 행이 있어 아무것도 저장하지 않았습니다.",
            code="IMPORT_ROWS_INVALID",
            details=errors,
        )
    counts = {"created": 0, "updated": 0}
    try:
        with transaction.atomic():
            for serializer in prepared:
                counts["updated" if serializer.instance else "created"] += 1
                serializer.save()
    except IntegrityError as exc:
        raise ValidationError(message="같은 분류·이름이 파일 안에 두 번 있습니다.") from exc
    return counts
